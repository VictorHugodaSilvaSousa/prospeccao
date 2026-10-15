"""Acha negócios locais no OpenStreetMap e marca quais não têm site.

Uso:
    python prospectar.py barbearia "Brasília"
    python prospectar.py shop=bakery "Taguatinga, DF"
Gera prospeccao-<nicho>-<cidade>.xlsx com os sem site primeiro.
"""
import html
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from google_places import buscar_google, e_rede_social
from verificar import conferir, normal

MODELO = Path(__file__).with_name("mensagem.txt")
MODELO_MELHORAR = Path(__file__).with_name("mensagem-melhorar.txt")
for _l in (Path(__file__).with_name(".env").read_text(encoding="utf-8").splitlines()
           if Path(__file__).with_name(".env").exists() else []):
    if "=" in _l and not _l.startswith("#"):  # chaves locais (GOOGLE_PLACES_KEY, BRAVE_API_KEY); fora do git
        os.environ.setdefault(*(p.strip() for p in _l.split("=", 1)))
UA ="prospeccao-local/1.0 (uso pessoal)"  # Nominatim exige User-Agent identificável
NICHOS = {
    "barbearia": '["shop"="hairdresser"]',
    "salao": '["shop"="hairdresser"]',
    "beleza": '["shop"="beauty"]',
    "padaria": '["shop"="bakery"]',
    "pizzaria": '["amenity"="restaurant"]["cuisine"~"pizza"]',
    "restaurante": '["amenity"="restaurant"]',
    "lanchonete": '["amenity"="fast_food"]',
    "cafe": '["amenity"="cafe"]',
    "academia": '["leisure"="fitness_centre"]',
    "dentista": '["amenity"="dentist"]',
    "clinica": '["amenity"="clinic"]',
    "petshop": '["shop"="pet"]',
    "veterinario": '["amenity"="veterinary"]',
    "oficina": '["shop"="car_repair"]',
    "roupas": '["shop"="clothes"]',
    "mercado": '["shop"="supermarket"]',
    "farmacia": '["amenity"="pharmacy"]',
}
SITE = ("website", "contact:website", "url")
SOCIAL = ("contact:instagram", "contact:facebook")
FONE = ("phone", "contact:phone", "contact:mobile", "contact:whatsapp")


def get(url, data=None):
    req = urllib.request.Request(url, data=data, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=180) as r:
        return json.load(r)


OVERPASS = ("https://overpass-api.de/api/interpreter", "https://overpass.private.coffee/api/interpreter",
            "https://overpass.kumi.systems/api/interpreter")


def overpass(data):
    """Servidores públicos vivem lotados (429/504): tenta os espelhos, 2 voltas com pausa."""
    erro = None
    for volta in range(2):
        for url in OVERPASS:
            try:
                return get(url, data)
            except (urllib.error.URLError, TimeoutError) as ex:
                erro = ex
                print(f"  {url.split('/')[2]} falhou ({ex}); tentando outro...")
        time.sleep(30)
    raise erro


def filtro(nicho):
    if nicho in NICHOS:
        return NICHOS[nicho]
    if re.fullmatch(r"[\w:]+=[\w;:-]+", nicho):
        k, v = nicho.split("=")
        return f'["{k}"="{v}"]'
    sys.exit(f"Nicho desconhecido: {nicho}. Use um de {', '.join(NICHOS)} ou chave=valor do OSM.")


def area_id(cidade):
    q = urllib.parse.urlencode({"q": cidade, "format": "json", "limit": 5, "countrycodes": "br"})
    for r in get(f"https://nominatim.openstreetmap.org/search?{q}"):
        if r["osm_type"] == "relation":
            return 3600000000 + int(r["osm_id"]), r["display_name"]
    sys.exit(f"Cidade não encontrada como área: {cidade}")


def linha(tags):
    """Converte as tags OSM de um negócio numa linha da planilha."""
    pega = lambda ks: next((tags[k] for k in ks if tags.get(k)), "")
    end = ", ".join(p for p in (
        " ".join(p for p in (tags.get("addr:street", ""), tags.get("addr:housenumber", "")) if p),
        tags.get("addr:suburb", ""),
    ) if p)
    site, social = pega(SITE), pega(SOCIAL)
    if site and e_rede_social(site):  # Instagram no campo "site" não é site próprio
        site, social = "", social or site
    return {
        "Nome": tags.get("name", ""),
        "Tem site": "sim" if site else "NÃO",
        "Site": site,
        "Rede social": social,
        "Telefone": pega(FONE),
        "Endereço": end,
        "Conferido": "",
    }


def celular(fones):
    """Primeiro celular BR (com 9 na frente) como 55DDD9XXXXXXXX; fixo não tem WhatsApp garantido."""
    for f in re.split(r"[;,/]", fones):
        d = re.sub(r"\D", "", f)
        d = d[2:] if d.startswith("55") and len(d) in (12, 13) else d.lstrip("0")
        if len(d) == 11 and d[2] == "9":
            return "55" + d
    return ""


def e_lead(l):
    """Vale abordar: tem celular e (não tem site OU o site tem problema)."""
    return bool(l["WhatsApp"]) and (l["Tem site"] == "NÃO" or bool(l.get("Problema")))


def ordem(l):
    return (not e_lead(l), l["Tem site"] == "sim", not l["Telefone"], l["Nome"])


def preparar(l, nicho):
    """Mensagem certa para o caso: sem site ou site com problema."""
    if l.get("Problema"):
        modelo = MODELO_MELHORAR.read_text(encoding="utf-8").strip()
    else:
        modelo = MODELO.read_text(encoding="utf-8").strip()
    cel = celular(l["Telefone"])
    l["Mensagem"] = modelo.format(nome=l["Nome"], nicho=nicho, problema=l.get("Problema", ""))
    l["WhatsApp"] = f"https://wa.me/{cel}?text={urllib.parse.quote(l['Mensagem'])}" if cel else ""


def conferir_todos(linhas, cidade, nicho=""):
    with ThreadPoolExecutor(6) as ex:  # cada conferência espera rede; em paralelo leva segundos, não minutos
        list(ex.map(lambda l: conferir(l, cidade), linhas))
    for l in linhas:
        preparar(l, nicho)


def juntar(google, osm):
    """Google primeiro (dados melhores); do OSM só entra quem o Google não trouxe (mesmo fone ou nome)."""
    fones = {re.sub(r"\D", "", l["Telefone"])[-8:] for l in google if l["Telefone"]}
    nomes = {"".join(normal(l["Nome"])) for l in google}
    extra = [l for l in osm if "".join(normal(l["Nome"])) not in nomes
             and not (l["Telefone"] and re.sub(r"\D", "", l["Telefone"])[-8:] in fones)]
    return google + extra


def buscar(nicho, cidade):
    """Retorna (nome da área, negócios ordenados com os que valem abordar primeiro)."""
    f = filtro(nicho)
    aid, nome = area_id(cidade)
    print(f"Buscando {nicho} em {nome}...")
    query = f"[out:json][timeout:170];area({aid})->.a;nwr{f}(area.a);out tags center;"
    data = urllib.parse.urlencode({"data": query}).encode()
    try:
        elementos = overpass(data)["elements"]
    except urllib.error.URLError:
        if not os.environ.get("GOOGLE_PLACES_KEY"):
            raise
        print("  OpenStreetMap fora do ar; seguindo só com o Google.")
        elementos = []
    osm = [dict(linha(e.get("tags", {})), Id=f"o:{e['type']}{e['id']}")
           for e in elementos if e.get("tags", {}).get("name")]
    google = buscar_google(f"{nicho} em {cidade}") if nicho in NICHOS else []
    if google:
        print(f"  Google: {len(google)} lugares; OpenStreetMap: {len(osm)}")
    linhas = juntar(google, osm)
    for l in linhas:
        l.setdefault("Fonte", "OpenStreetMap")
        l["Status"], l["Problema"] = "", ""
        preparar(l, nicho)
    linhas.sort(key=ordem)
    return nome, linhas


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    nicho, cidade = sys.argv[1].lower(), sys.argv[2]
    nome, linhas = buscar(nicho, cidade)
    alvo = [l for l in linhas if l["WhatsApp"]]
    print(f"Conferindo o site de {len(alvo)} negócios com celular...")
    conferir_todos(alvo, cidade, nicho)
    linhas.sort(key=ordem)

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospecção"
    cols = ["Nome", "Tem site", "Problema", "WhatsApp", "Status", "Telefone", "Rede social", "Site", "Endereço",
            "Conferido", "Fonte", "Mensagem"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for l in linhas:
        ws.append([l[c] for c in cols])
        if l["WhatsApp"]:
            c = ws.cell(ws.max_row, 4, "Abrir conversa")
            c.hyperlink, c.style = l["WhatsApp"], "Hyperlink"
    for col, w in zip("ABCDEFGHIJKL", (40, 10, 30, 16, 14, 20, 30, 35, 45, 16, 14, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    slug = re.sub(r"\W+", "-", f"{nicho}-{cidade}".lower()).strip("-")
    arq = f"prospeccao-{slug}.xlsx"
    wb.save(arq)
    sem = sum(l["Tem site"] == "NÃO" for l in linhas)
    ruim = sum(bool(l["Problema"]) for l in linhas)
    leads = sum(e_lead(l) for l in linhas)
    print(f"{len(linhas)} negócios, {sem} sem site, {ruim} com site com problema; {leads} para abordar. Planilha: {arq}")
    pagina = Path(arq).with_suffix(".html").resolve()
    pagina.write_text(html_tabela(f"{nicho} — {nome}", linhas), encoding="utf-8")
    print(f"Página: {pagina.name}")
    webbrowser.open(pagina.as_uri())  # o PC pode não ter Excel; navegador sempre tem


def html_tabela(titulo, linhas):
    e = html.escape
    zap = lambda u: f'<a href="{e(u)}" target="_blank">WhatsApp</a>' if u else ""
    tr = "".join(
        f"<tr><td>{e(l['Nome'])}</td><td>{e(l['Tem site'])}</td><td>{e(l['Problema'] or l['Site'])}</td>"
        f"<td>{zap(l['WhatsApp']) if e_lead(l) else ''}</td>"
        f"<td>{e(l['Telefone'])}</td><td>{e(l['Rede social'])}</td><td>{e(l['Endereço'])}</td></tr>"
        for l in linhas
    )
    return f"""<!doctype html><meta charset="utf-8"><title>{e(titulo)}</title>
<style>body{{font-family:sans-serif;margin:16px}}table{{border-collapse:collapse;width:100%}}
td,th{{border:1px solid #ccc;padding:6px;text-align:left}}th{{background:#eee;position:sticky;top:0}}
a{{background:#25d366;color:#fff;padding:4px 10px;border-radius:4px;text-decoration:none}}</style>
<h2>{e(titulo)}</h2><p>{len(linhas)} negócios. Clique em WhatsApp para abrir a conversa com a mensagem pronta.</p>
<table><tr><th>Nome</th><th>Tem site</th><th>Site / problema</th><th>Contato</th><th>Telefone</th><th>Rede social</th><th>Endereço</th></tr>{tr}</table>"""


if __name__ == "__main__":
    main()
